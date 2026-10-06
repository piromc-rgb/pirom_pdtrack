import os
import re
import json
from datetime import datetime
import openpyxl

# Google Drive folder: https://drive.google.com/drive/folders/1z4qVl5Iikwd1PTQSdwo0Et_nIk9VMTGS
GDRIVE_DIR = r'G:\My Drive\Purchase Data'
FILE_KEYWORD = 'Report PO'


def parse_file_date(filename, filepath):
    """Extract dd-mm-yy from names like 'Report PO ค้างรับ  06-10-26.xlsx'; fall back to mtime."""
    m = re.search(r'(\d{1,2})-(\d{1,2})-(\d{2,4})', filename)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y >= 2500:
            y -= 543
        elif y < 100:
            y += 2000
        try:
            return datetime(y, mo, d)
        except ValueError:
            pass
    return datetime.fromtimestamp(os.path.getmtime(filepath))


def find_latest_file():
    files = []
    for f in os.listdir(GDRIVE_DIR):
        if f.lower().endswith('.xlsx') and not f.startswith('~$') and FILE_KEYWORD.lower() in f.lower():
            p = os.path.join(GDRIVE_DIR, f)
            files.append((parse_file_date(f, p), os.path.getmtime(p), p, f))
    if not files:
        raise FileNotFoundError(f'No "{FILE_KEYWORD}" xlsx files in {GDRIVE_DIR}')
    files.sort(reverse=True)
    return files[0]


def fmt(v):
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y')
    return ''


def main():
    dt, _, path, name = find_latest_file()
    print(f'Latest file: {name} ({dt:%Y-%m-%d})')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Data'] if 'Data' in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else '' for h in next(rows)]
    col = {h: i for i, h in enumerate(headers) if h}
    order_i, line_i = col['Order'], col['Line']
    # the item code sits in the column right after the (empty) 'Item' column
    item_i = col['Item'] + 1
    qty_i = col['Ordered Quantity']
    plan_i, conf_i = col['Planned Receipt Date'], col['Confirmed Receipt Date']

    lines = []
    for r in rows:
        order = str(r[order_i] or '').strip().upper()
        if not order:
            continue
        lines.append({
            'order': order,
            'line': r[line_i],
            'item': str(r[item_i] or '').strip().upper(),
            'qty': r[qty_i],
            'planned': fmt(r[plan_i]),
            'confirmed': fmt(r[conf_i]),
        })

    out = {'fileName': name, 'fileDate': dt.strftime('%Y-%m-%d'), 'lines': lines}
    target = os.path.join(os.path.dirname(__file__), '..', 'src', 'data', 'poPending.json')
    with open(target, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False)
    print(f'Wrote {len(lines)} pending PO lines -> {os.path.abspath(target)}')


if __name__ == '__main__':
    main()
