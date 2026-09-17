import os
import re
import json
import shutil
from datetime import datetime
import openpyxl

# Google Drive folder link: https://drive.google.com/drive/folders/1Yt8drFmq0END9fAEWUy0No6sZ76H1dtA
GDRIVE_DIR = r'G:\My Drive\staus overview'

def parse_file_date(filename, filepath=None):
    """
    Extract date from filename supporting patterns:
    - 69-09-17 (Thai BE year 2-digit e.g. 2569 -> 2026)
    - 26-09-15 (CE year 2-digit e.g. 2026)
    - 2026-09-15 (CE year 4-digit)
    - 2569-09-17 (BE year 4-digit)
    """
    m = re.search(r'(\d{2,4})-(\d{2})-(\d{2})', filename)
    if m:
        y, month, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y >= 2500:
            y -= 543
        elif y >= 50 and y <= 99:
            y = 1957 + y
        elif y < 50:
            y = 2000 + y
        try:
            return datetime(y, month, d)
        except:
            pass
    if filepath and os.path.exists(filepath):
        return datetime.fromtimestamp(os.path.getmtime(filepath))
    return datetime.min

def find_latest_status_overview_file():
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    
    # 1. First priority: Target Google Drive folder (1Yt8drFmq0END9fAEWUy0No6sZ76H1dtA)
    gdrive_files = []
    if os.path.exists(GDRIVE_DIR):
        print(f"Checking Google Drive folder: {GDRIVE_DIR}")
        for f in os.listdir(GDRIVE_DIR):
            if f.endswith('.xlsx') and not f.startswith('~$'):
                full_p = os.path.join(GDRIVE_DIR, f)
                dt = parse_file_date(f, full_p)
                mtime = os.path.getmtime(full_p)
                gdrive_files.append((dt, mtime, full_p, f, 'Google Drive'))
                
        if gdrive_files:
            # Sort by date descending, then mtime descending
            gdrive_files.sort(key=lambda x: (x[0], x[1]), reverse=True)
            print(f"Found {len(gdrive_files)} file(s) in Google Drive folder.")
            return gdrive_files[0]

    # 2. Secondary fallback: Local project and backup folders
    print("Checking local fallback directories...")
    search_dirs = [
        base_dir,
        os.path.abspath(os.path.join(base_dir, '..', 'AMW Packing List')),
        os.path.abspath(os.path.join(base_dir, '..', 'PD Plan copy')),
        os.path.abspath(os.path.join(base_dir, '..', 'PD Plan')),
        os.path.abspath(os.path.join(base_dir, '..', 'PD_Tracking')),
    ]
    
    local_files = []
    seen = set()
    for d in search_dirs:
        if not os.path.exists(d):
            continue
        try:
            for f in os.listdir(d):
                if f.endswith('.xlsx') and 'Status Overview' in f and not f.startswith('~$'):
                    full_p = os.path.join(d, f)
                    if full_p in seen:
                        continue
                    seen.add(full_p)
                    dt = parse_file_date(f, full_p)
                    mtime = os.path.getmtime(full_p)
                    local_files.append((dt, mtime, full_p, f, 'Local'))
        except Exception as e:
            print(f"Warning scanning {d}: {e}")
            
    if not local_files:
        raise FileNotFoundError("No Status Overview xlsx files found in Google Drive or local folders.")
        
    local_files.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return local_files[0]

def process_overview_file(filepath):
    print(f"Loading workbook: {filepath} ...")
    wb = openpyxl.load_workbook(filepath, read_only=True)
    
    # Locate data sheet: strict priority to exact sheet named 'data' (case-insensitive)
    target_sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == 'data':
            target_sheet = name
            break
            
    # Secondary fallback if no exact 'data' sheet exists
    if not target_sheet:
        for name in wb.sheetnames:
            if name.strip().lower().startswith('data'):
                target_sheet = name
                break
                
    if not target_sheet:
        for name in wb.sheetnames:
            if 'data' in name.strip().lower():
                target_sheet = name
                break
                
    if not target_sheet:
        target_sheet = wb.sheetnames[0]
        
    print(f"Using sheet: '{target_sheet}'")
    ws = wb[target_sheet]
    
    headers = []
    for row in ws.iter_rows(max_row=1, values_only=True):
        headers = [str(x).strip() if x is not None else '' for x in row]
        break
    
    def find_col(keywords):
        # 1. Exact match first
        for i, h in enumerate(headers):
            if any(k.strip().lower() == h.strip().lower() for k in keywords):
                return i
        # 2. Substring match
        for i, h in enumerate(headers):
            if any(k.strip().lower() in h.strip().lower() for k in keywords):
                return i
        return -1
    
    pd_idx = find_col(['Production Order', 'Prod Order', 'PD No', 'PD'])
    item_idx = find_col(['Item_5', 'เลขที่ Item', 'Item Code', 'Item No', 'itemCode'])
    desc_idx = find_col(['Description', 'ชื่อ Item', 'Item Name', 'itemName'])
    proj_idx = find_col(['Project', 'เลขที่โครงการ', 'Project No', 'projectCode'])
    cust_idx = find_col(['Customer', 'ลูกค้า', 'customer'])
    op_idx = find_col(['Operation', 'ลำดับ'])
    wc_idx = find_col(['Work Center', 'WorkCenter', 'WC', 'เครื่อง'])
    opdesc_idx = find_col(['r.ref.oper.desc', 'Operation Description', 'ชื่อ Operation'])
    opstatus_idx = find_col(['Operation Status', 'สถานะ Operation'])
    orderstatus_idx = find_col(['Order Status', 'Status', 'สถานะ'])
    
    print(f"Column indexes: PD={pd_idx}, Item={item_idx}, Op={op_idx}, OpDesc={opdesc_idx}, OpStatus={opstatus_idx}, OrderStatus={orderstatus_idx}")
    
    orders = {}
    row_count = 0
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_count += 1
        pd_val = str(row[pd_idx] or '').strip().upper() if pd_idx != -1 and pd_idx < len(row) else ''
        item_val = str(row[item_idx] or '').strip().upper() if item_idx != -1 and item_idx < len(row) else ''
        if not pd_val and not item_val:
            continue
            
        desc_val = str(row[desc_idx] or '').strip() if desc_idx != -1 and desc_idx < len(row) else ''
        proj_val = str(row[proj_idx] or '').strip().upper() if proj_idx != -1 and proj_idx < len(row) else ''
        cust_val = str(row[cust_idx] or '').strip() if cust_idx != -1 and cust_idx < len(row) else ''
        op_val = row[op_idx] if op_idx != -1 and op_idx < len(row) else ''
        wc_val = str(row[wc_idx] or '').strip() if wc_idx != -1 and wc_idx < len(row) else ''
        opdesc_val = str(row[opdesc_idx] or '').strip() if opdesc_idx != -1 and opdesc_idx < len(row) else ''
        opstatus_val = str(row[opstatus_idx] or '').strip() if opstatus_idx != -1 and opstatus_idx < len(row) else ''
        orderstatus_val = str(row[orderstatus_idx] or '').strip() if orderstatus_idx != -1 and orderstatus_idx < len(row) else ''
        
        # Clean up op number format (e.g. 10.0 -> 10)
        clean_op = op_val
        op_num = 999
        try:
            op_float = float(str(op_val).strip())
            op_num = int(op_float)
            clean_op = str(op_num)
        except:
            pass
            
        if pd_val:
            if pd_val not in orders:
                orders[pd_val] = {
                    'pd': pd_val,
                    'itemCode': item_val,
                    'itemName': desc_val,
                    'project': proj_val,
                    'customer': cust_val,
                    'orderStatus': orderstatus_val,
                    'ops': []
                }
            if opdesc_val or op_val or opstatus_val:
                orders[pd_val]['ops'].append({
                    'op': clean_op,
                    'opNum': op_num,
                    'wc': wc_val,
                    'desc': opdesc_val,
                    'opStatus': opstatus_val,
                    'orderStatus': orderstatus_val
                })
    
    print(f"Read {row_count} rows, found {len(orders)} unique Production Orders.")
    
    # Process each order
    byPd = {}
    byItem = {}
    byProjItem = {}
    
    for pd_val, data in orders.items():
        ops = sorted(data['ops'], key=lambda x: x['opNum'])
        
        completed_ops = [o for o in ops if o['opStatus'].lower() in ['completed', 'close', 'closed', 'เสร็จแล้ว', 'เสร็จสิ้น']]
        active_ops = [o for o in ops if o['opStatus'].lower() == 'active']
        ready_ops = [o for o in ops if o['opStatus'].lower() == 'ready to start']
        pending_ops = [o for o in ops if o['opStatus'].lower() not in ['completed', 'close', 'closed', 'เสร็จแล้ว', 'เสร็จสิ้น', 'active']]
        
        all_completed = len(ops) > 0 and len(completed_ops) == len(ops)
        last_completed_op = completed_ops[-1] if completed_ops else None
        active_op = active_ops[0] if active_ops else None
        ready_op = ready_ops[0] if ready_ops else (pending_ops[0] if pending_ops else None)
        
        # Construct meta
        final_status = data['orderStatus'] or ''
        meta = {
            'status': final_status,
            'customer': data['customer'],
            'project': data['project'],
            'itemCode': data['itemCode'],
            'description': data['itemName'],
            'prodOrder': pd_val,
        }
        
        # Record latest completed operation if any
        if last_completed_op:
            meta['lastCompletedOp'] = f"Op {last_completed_op['op']}"
            meta['lastCompletedOpDesc'] = last_completed_op['desc']
            meta['lastCompletedOpWc'] = last_completed_op['wc']
            meta['lastCompletedOpNo'] = last_completed_op['op']
        
        if all_completed or (data['orderStatus'] and data['orderStatus'].lower() in ['closed', 'completed', 'เสร็จแล้ว', 'เสร็จสิ้น']):
            meta['status'] = 'Closed'
            meta['isAllCompleted'] = True
            meta['currentOp'] = 'เสร็จทุกขั้นตอน'
            meta['currentOpDesc'] = 'Completed'
            meta['currentOpStatus'] = 'Completed'
        elif active_op:
            wc = active_op['wc']
            label = f"Op {active_op['op']}: {active_op['desc']}{f' ({wc})' if wc else ''}"
            meta['activeOp'] = label
            meta['activeOpDesc'] = active_op['desc']
            meta['activeOpWc'] = wc
            meta['activeOpNo'] = active_op['op']
            meta['currentOp'] = label
            meta['currentOpDesc'] = active_op['desc']
            meta['currentOpStatus'] = 'Active'
            meta['status'] = 'Active'
        elif ready_op:
            # The next pending op is "รอขึ้น" (waiting to start)
            wc = ready_op['wc']
            label = f"Op {ready_op['op']}: {ready_op['desc']}{f' ({wc})' if wc else ''}"
            meta['readyOp'] = label
            meta['readyOpDesc'] = ready_op['desc']
            meta['readyOpWc'] = wc
            meta['readyOpNo'] = ready_op['op']
            meta['hasReadyOp'] = True
            meta['currentOp'] = label
            meta['currentOpDesc'] = ready_op['desc']
            meta['currentOpStatus'] = 'Ready to Start' if ready_op['opStatus'].lower() == 'ready to start' else 'Planned'
            meta['status'] = 'Ready to Start' if ready_op['opStatus'].lower() == 'ready to start' else 'Planned'
            
        byPd[pd_val] = meta
        
        # Index in byProjItem and byItem
        item_code = data['itemCode']
        proj = data['project']
        if item_code:
            if item_code not in byItem or meta['status'] == 'Active':
                byItem[item_code] = meta
            if proj:
                proj_key = f"{proj}|{item_code}"
                if proj_key not in byProjItem or meta['status'] == 'Active':
                    byProjItem[proj_key] = meta
                    
    return byPd, byItem, byProjItem

def main():
    print("=== Auto-detecting Latest Status Overview File ===")
    dt, mtime, filepath, filename, source = find_latest_status_overview_file()
    print(f"\n>>> Selected LATEST file from [{source}]: {filename} ({dt.strftime('%Y-%m-%d')})")
    print(f"Path: {filepath}")
    
    # Copy to project root if not already there
    project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    local_target = os.path.join(project_root, filename)
    if os.path.abspath(filepath) != os.path.abspath(local_target):
        shutil.copy2(filepath, local_target)
        print(f"Copied latest file to project root: {local_target}")
        
    byPd, byItem, byProjItem = process_overview_file(filepath)
    
    # Save to src/data
    status_data_path = os.path.join(project_root, 'src', 'data', 'overviewStatusData.json')
    item_map_path = os.path.join(project_root, 'src', 'data', 'overviewItemMap.json')
    
    print(f"Writing {status_data_path} ...")
    with open(status_data_path, 'w', encoding='utf-8') as f:
        json.dump(byPd, f, ensure_ascii=False, indent=2)
        
    print(f"Writing {item_map_path} ...")
    with open(item_map_path, 'w', encoding='utf-8') as f:
        json.dump({'byItem': byItem, 'byProjItem': byProjItem}, f, ensure_ascii=False, indent=2)
        
    print("\n=== Verification ===")
    for target_pd in ['PD2607680', 'PD2608432']:
        meta = byPd.get(target_pd)
        if meta:
            print(f"{target_pd}:")
            print(f"  Status: {meta.get('status')}")
            print(f"  activeOp: {meta.get('activeOp')}")
            print(f"  activeOpDesc: {meta.get('activeOpDesc')}")
            print(f"  readyOp: {meta.get('readyOp')}")
            print(f"  readyOpDesc: {meta.get('readyOpDesc')}")
            print(f"  currentOp: {meta.get('currentOp')}")
        else:
            print(f"{target_pd}: NOT FOUND")

if __name__ == '__main__':
    main()
